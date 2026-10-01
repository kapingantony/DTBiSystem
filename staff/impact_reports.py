"""Excel and PDF exports for the public four-year startup impact explorer."""

import io
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList


EXCEL_MIME = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def _report_title(period_label):
    return f'DTBi / BUNI Startup Impact Summary ({period_label})'


def startup_impact_xlsx(period_label, summary, annual_rows, startup_rows, as_of_label):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Summary'
    sheet.merge_cells('A1:B1')
    sheet['A1'] = _report_title(period_label)
    sheet['A1'].font = Font(bold=True, size=16, color='FFFFFF')
    sheet['A1'].fill = PatternFill('solid', fgColor='12304D')
    sheet['A1'].alignment = Alignment(vertical='center')
    sheet.row_dimensions[1].height = 28
    sheet.append(['Measure', 'Value'])
    measures = [
        ('Published startups in cohort', 'published_startups'),
        ('Participant journeys started', 'participant_journeys_started'),
        ('Participant support activities', 'support_activities'),
        ('Participants receiving support', 'participants_supported'),
        ('Completed mentoring sessions', 'mentor_sessions'),
        ('Participants with outcome records', 'participants_with_outcomes'),
        ('Participant outcome snapshots', 'outcome_snapshots'),
        ('Reported full-time jobs', 'full_time_jobs'),
        ('Reported part-time jobs', 'part_time_jobs'),
    ]
    for row in measures:
        sheet.append([row[0], summary[row[1]]])
    sheet.append([])
    sheet.append(['Reporting period', f'Activity is grouped by startup cohort year. Report prepared {as_of_label}.'])
    _style_table(sheet, header_row=2)
    sheet.column_dimensions['A'].width = 40
    sheet.column_dimensions['B'].width = 86

    charts = workbook.create_sheet('Charts')
    charts.append(['Year', 'Journeys', 'Support', 'Mentoring', 'Outcomes'])
    for row in annual_rows:
        charts.append([row['year'], row['participant_journeys_started'], row['support_activities'], row['mentor_sessions'], row['participants_with_outcomes']])
    _style_table(charts, header_row=1)
    if annual_rows:
        activity_chart = BarChart()
        activity_chart.type = 'col'
        activity_chart.style = 10
        activity_chart.title = 'Annual programme activity'
        activity_chart.y_axis.title = 'People and activities'
        activity_chart.x_axis.title = 'Year'
        activity_chart.grouping = 'clustered'
        activity_chart.overlap = 0
        activity_chart.add_data(Reference(charts, min_col=2, max_col=5, min_row=1, max_row=len(annual_rows) + 1), titles_from_data=True)
        activity_chart.set_categories(Reference(charts, min_col=1, min_row=2, max_row=len(annual_rows) + 1))
        activity_chart.height = 9
        activity_chart.width = 18
        charts.add_chart(activity_chart, 'A8')
    charts['G1'] = 'Employment type'
    charts['H1'] = 'Jobs'
    charts['G2'] = 'Full-time'
    charts['H2'] = summary['full_time_jobs']
    charts['G3'] = 'Part-time'
    charts['H3'] = summary['part_time_jobs']
    employment_chart = PieChart()
    employment_chart.title = 'Reported jobs by type'
    employment_chart.style = 10
    employment_chart.add_data(Reference(charts, min_col=8, min_row=1, max_row=3), titles_from_data=True)
    employment_chart.set_categories(Reference(charts, min_col=7, min_row=2, max_row=3))
    employment_chart.dataLabels = DataLabelList()
    employment_chart.dataLabels.showPercent = True
    employment_chart.height = 8
    employment_chart.width = 12
    charts.add_chart(employment_chart, 'G8')
    charts.column_dimensions['A'].width = 12
    charts.column_dimensions['G'].width = 20
    charts.column_dimensions['H'].width = 14

    annual = workbook.create_sheet('Annual breakdown')
    annual.append(['Year', 'Published startups in cohort', 'Journeys started', 'Support activities', 'Participants supported', 'Completed mentor sessions', 'Participants with outcomes', 'Outcome snapshots', 'Full-time jobs', 'Part-time jobs'])
    for row in annual_rows:
        annual.append([row['year'], row['published_startups'], row['participant_journeys_started'], row['support_activities'], row['participants_supported'], row['mentor_sessions'], row['participants_with_outcomes'], row['outcome_snapshots'], row['full_time_jobs'], row['part_time_jobs']])
    _style_table(annual, header_row=1)
    _fit_columns(annual, maximum=36)

    startups = workbook.create_sheet('Startup records')
    startups.append(['Startup organization', 'Category', 'Cohort year', 'Journeys started', 'Support activities', 'Participants supported', 'Completed mentor sessions', 'Participants with outcomes', 'Outcome snapshots', 'Full-time jobs', 'Part-time jobs'])
    for row in startup_rows:
        startups.append([row['startup'].name, row['startup'].industry or 'Not recorded', row['cohort_year'] or 'Not recorded', row['participant_journeys_started'], row['support_activities'], row['participants_supported'], row['mentor_sessions'], row['participants_with_outcomes'], row['outcome_snapshots'], row['full_time_jobs'], row['part_time_jobs']])
    _style_table(startups, header_row=1)
    _fit_columns(startups, maximum=42)

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _style_table(sheet, header_row):
    for cell in sheet[header_row]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='176C78')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    sheet.freeze_panes = f'A{header_row + 1}'
    sheet.auto_filter.ref = sheet.dimensions
    for row in sheet.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)


def _fit_columns(sheet, maximum=42):
    for column in sheet.columns:
        letter = get_column_letter(column[0].column)
        width = max((len(str(cell.value or '')) for cell in column), default=10) + 2
        sheet.column_dimensions[letter].width = min(max(width, 12), maximum)


def startup_impact_pdf(period_label, summary, annual_rows, startup_rows, as_of_label):
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.graphics.shapes import Drawing, String
        from reportlab.graphics.charts.piecharts import Pie
        from reportlab.graphics.charts.barcharts import VerticalBarChart
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError as exc:
        raise RuntimeError('PDF export requires the reportlab package declared in requirements.txt.') from exc

    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=landscape(A4), rightMargin=14 * mm, leftMargin=14 * mm, topMargin=14 * mm, bottomMargin=14 * mm, title=_report_title(period_label))
    styles = getSampleStyleSheet()
    story = [
        Paragraph(_report_title(period_label), styles['Title']),
        Paragraph('Published startup cohort and programme impact figures', styles['Heading2']),
        Paragraph(f'Counts include published startups whose cohort year falls within the selected four-year range. Activity is grouped by startup cohort year. Report prepared {escape(as_of_label)}.', styles['BodyText']),
        Spacer(1, 8 * mm),
    ]
    chart_drawing = Drawing(720, 205)
    chart_drawing.add(String(12, 190, 'Reported jobs by type', fontSize=11, fillColor=colors.HexColor('#12304D')))
    if summary['full_time_jobs'] + summary['part_time_jobs']:
        pie = Pie()
        pie.x = 18
        pie.y = 22
        pie.width = 145
        pie.height = 145
        pie.data = [summary['full_time_jobs'], summary['part_time_jobs']]
        pie.labels = ['Full-time', 'Part-time']
        pie.slices[0].fillColor = colors.HexColor('#176C78')
        pie.slices[1].fillColor = colors.HexColor('#4C8DF6')
        pie.slices.strokeWidth = 0.5
        chart_drawing.add(pie)
        chart_drawing.add(String(170, 120, f"Full-time: {summary['full_time_jobs']}", fontSize=9, fillColor=colors.HexColor('#176C78')))
        chart_drawing.add(String(170, 101, f"Part-time: {summary['part_time_jobs']}", fontSize=9, fillColor=colors.HexColor('#4C8DF6')))
    else:
        chart_drawing.add(String(24, 100, 'No reported jobs in this period', fontSize=9, fillColor=colors.HexColor('#617589')))
    chart_drawing.add(String(330, 190, 'Annual programme activity', fontSize=11, fillColor=colors.HexColor('#12304D')))
    if annual_rows:
        annual_chart = VerticalBarChart()
        annual_chart.x = 330
        annual_chart.y = 34
        annual_chart.width = 365
        annual_chart.height = 135
        annual_chart.data = [
            [row['participant_journeys_started'] for row in annual_rows],
            [row['support_activities'] for row in annual_rows],
            [row['mentor_sessions'] for row in annual_rows],
            [row['participants_with_outcomes'] for row in annual_rows],
        ]
        annual_chart.categoryAxis.categoryNames = [str(row['year']) for row in annual_rows]
        annual_chart.categoryAxis.labels.fontSize = 7
        annual_chart.valueAxis.labels.fontSize = 7
        annual_chart.groupSpacing = 8
        annual_chart.barSpacing = 2
        for index, color in enumerate(['#176C78', '#4C8DF6', '#E0A53B', '#8762C5']):
            annual_chart.bars[index].fillColor = colors.HexColor(color)
        chart_drawing.add(annual_chart)
        for index, (label, color) in enumerate(zip(['Journeys', 'Support', 'Mentoring', 'Outcomes'], ['#176C78', '#4C8DF6', '#E0A53B', '#8762C5'])):
            x = 338 + index * 88
            chart_drawing.add(String(x, 17, label, fontSize=7, fillColor=colors.HexColor(color)))
    story.extend([chart_drawing, Spacer(1, 5 * mm)])
    summary_rows = [['Measure', 'Value']]
    for label, key in [
        ('Published startups in cohort', 'published_startups'),
        ('Participant journeys started', 'participant_journeys_started'),
        ('Participant support activities', 'support_activities'),
        ('Participants receiving support', 'participants_supported'),
        ('Completed mentoring sessions', 'mentor_sessions'),
        ('Participants with outcome records', 'participants_with_outcomes'),
        ('Participant outcome snapshots', 'outcome_snapshots'),
        ('Reported full-time jobs', 'full_time_jobs'),
        ('Reported part-time jobs', 'part_time_jobs'),
    ]:
        summary_rows.append([label, str(summary[key])])
    story.extend([Paragraph('Four-year summary', styles['Heading2']), Table(summary_rows, colWidths=[112 * mm, 34 * mm], repeatRows=1), Spacer(1, 7 * mm)])

    annual_rows_data = [['Year', 'Published', 'Journeys', 'Support activities', 'Supported people', 'Mentor sessions', 'Outcome people', 'Snapshots', 'Full-time jobs', 'Part-time jobs']]
    for row in annual_rows:
        annual_rows_data.append([str(row['year']), str(row['published_startups']), str(row['participant_journeys_started']), str(row['support_activities']), str(row['participants_supported']), str(row['mentor_sessions']), str(row['participants_with_outcomes']), str(row['outcome_snapshots']), str(row['full_time_jobs']), str(row['part_time_jobs'])])
    story.extend([Paragraph('Annual breakdown', styles['Heading2']), Table(annual_rows_data, colWidths=[17 * mm, 23 * mm, 22 * mm, 26 * mm, 26 * mm, 24 * mm, 27 * mm, 20 * mm, 22 * mm, 22 * mm], repeatRows=1), Spacer(1, 7 * mm)])

    story.append(Paragraph('Startup cohort detail', styles['Heading2']))
    startup_rows_data = [['Startup', 'Category', 'Cohort', 'Journeys', 'Support', 'Supported', 'Mentoring', 'Outcomes', 'Snapshots', 'Full-time jobs', 'Part-time jobs']]
    for row in startup_rows:
        startup_rows_data.append([Paragraph(escape(row['startup'].name), styles['BodyText']), Paragraph(escape(row['startup'].industry or 'Not recorded'), styles['BodyText']), str(row['cohort_year'] or '—'), str(row['participant_journeys_started']), str(row['support_activities']), str(row['participants_supported']), str(row['mentor_sessions']), str(row['participants_with_outcomes']), str(row['outcome_snapshots']), str(row['full_time_jobs']), str(row['part_time_jobs'])])
    story.append(Table(startup_rows_data, colWidths=[48 * mm, 32 * mm, 18 * mm, 18 * mm, 18 * mm, 20 * mm, 20 * mm, 20 * mm, 18 * mm, 22 * mm, 22 * mm], repeatRows=1))

    def style_report_table(table, compact=False):
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#12304D')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('GRID', (0, 0), (-1, -1), 0.35, colors.HexColor('#D3DEE7')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F2F6F8')]),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('FONTSIZE', (0, 0), (-1, -1), 7 if compact else 9),
            ('LEADING', (0, 0), (-1, -1), 9 if compact else 12),
            ('LEFTPADDING', (0, 0), (-1, -1), 5),
            ('RIGHTPADDING', (0, 0), (-1, -1), 5),
            ('TOPPADDING', (0, 0), (-1, -1), 5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ]))

    # Apply the same restrained DTBi header and table styling to each section.
    for index, flowable in enumerate(story):
        if isinstance(flowable, Table):
            style_report_table(flowable, compact=(len(flowable._cellvalues[0]) > 4))
    doc.build(story)
    return output.getvalue()
