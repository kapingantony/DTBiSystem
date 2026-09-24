import csv
import re
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.core.exceptions import ValidationError
from django.db import transaction
from django.db.models import Count, Sum
from django.utils import timezone
from openpyxl import load_workbook

from .models import (
    DataImportBatch, Funding, Investor, KPI, Mentor, MentorEngagement,
    PageVisit, Startup, StartupStatusHistory,
)


IMPORT_MODELS = {'startup': Startup, 'mentor': Mentor, 'investor': Investor}
IMPORT_FIELDS = {
    'startup': ('name', 'startup_type', 'description', 'industry', 'website', 'contact_email', 'phone', 'source', 'status', 'contract_status', 'founded_date', 'incubation_start', 'incubation_end', 'year_incubated'),
    'mentor': ('name', 'email', 'gender', 'education_level', 'phone', 'role', 'skills', 'training_topics', 'is_active', 'source'),
    'investor': ('name', 'organization', 'email', 'phone', 'website', 'description', 'investment_interest', 'status', 'source'),
}
ALIASES = {
    'name': ('name', 'startup', 'startup name', 'company', 'company name', 'mentor name', 'investor name', 'full name', 'combined full name'),
    'organization': ('organization', 'organisation', 'company', 'company name', 'fund'),
    'email': ('email', 'email address', 'contact email'),
    'contact_email': ('contact email', 'email', 'email address'),
    'industry': ('industry', 'sector'),
    'website': ('website', 'url', 'web address'),
    'phone': ('phone', 'telephone', 'mobile', 'contact number'),
    'description': ('description', 'profile', 'about', 'summary'),
    'status': ('status', 'startup status', 'investor status'),
    'startup_type': ('startup type', 'type'),
    'contract_status': ('contract status', 'contract'),
    'founded_date': ('founded date', 'date founded', 'year founded'),
    'incubation_start': ('incubation start', 'programme start', 'program start'),
    'incubation_end': ('incubation end', 'programme end', 'program end'),
    'year_incubated': ('year incubated', 'cohort year', 'incubation year'),
    'gender': ('gender',),
    'education_level': ('education', 'education level', 'qualification'),
    'role': ('role', 'mentor role', 'describe yourself'),
    'skills': ('skills', 'skill set', 'expertise'),
    'training_topics': ('training topics', 'topics', 'areas to train'),
    'investment_interest': ('investment interest', 'investment focus', 'interests'),
    'source': ('source', 'data source', 'programme'),
    'is_active': ('active', 'is active', 'status'),
}


def _normalize(value):
    return re.sub(r'[^a-z0-9]+', ' ', str(value or '').strip().lower()).strip()


def read_import_file(path):
    """Return column headings and non-empty rows for CSV, XLSX, or tabular PDF."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == '.csv':
        try:
            handle = path.open('r', encoding='utf-8-sig', newline='')
            rows = list(csv.reader(handle))
        except UnicodeDecodeError:
            with path.open('r', encoding='cp1252', newline='') as handle:
                rows = list(csv.reader(handle))
        finally:
            if 'handle' in locals():
                handle.close()
    elif suffix == '.xlsx':
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = list(workbook.active.iter_rows(values_only=True))
        finally:
            workbook.close()
    elif suffix == '.pdf':
        try:
            import pdfplumber
        except ImportError as exc:
            raise ValidationError('PDF importing needs the declared pdfplumber dependency. Restart the app and install requirements.txt, then retry.') from exc
        rows = []
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                rows.extend(row for table in (page.extract_tables() or []) for row in table if row)
        if not rows:
            raise ValidationError('No tables could be read from this PDF. Use a text based PDF with a clear table, or convert it to Excel/CSV.')
    else:
        raise ValidationError('Upload an .xlsx, .csv, or table based .pdf file.')

    rows = [list(row) for row in rows if row and any(value not in (None, '') for value in row)]
    if len(rows) < 2:
        raise ValidationError('The file must contain a heading row and at least one data row.')
    headers = [str(value or '').strip() for value in rows[0]]
    if not any(headers):
        raise ValidationError('The first row must contain column headings.')
    records = []
    for row in rows[1:20001]:
        records.append({headers[index]: value for index, value in enumerate(row) if index < len(headers) and headers[index]})
    if not records:
        raise ValidationError('No data rows were found.')
    return headers, records


def guess_field_mapping(dataset, headers):
    normalized = {_normalize(header): header for header in headers}
    mapping = {}
    for field in IMPORT_FIELDS[dataset]:
        aliases = ALIASES.get(field, (field.replace('_', ' '),))
        found = next((normalized[_normalize(alias)] for alias in aliases if _normalize(alias) in normalized), None)
        if found:
            mapping[field] = found
    return mapping


def _value_for_field(raw, field):
    if raw is None:
        return None
    if isinstance(raw, str):
        raw = raw.strip()
    if raw == '':
        return None
    if field in {'founded_date', 'incubation_start', 'incubation_end'}:
        if isinstance(raw, datetime):
            return raw.date()
        if isinstance(raw, date):
            return raw
        value = str(raw).strip()
        for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%m/%d/%Y', '%Y'):
            try:
                return datetime.strptime(value, fmt).date()
            except ValueError:
                continue
        raise ValidationError(f'Invalid date: {raw}')
    if field == 'year_incubated':
        try:
            return int(float(raw))
        except (TypeError, ValueError):
            raise ValidationError(f'Invalid year: {raw}')
    if field == 'is_active':
        value = _normalize(raw)
        if value in {'yes', 'true', '1', 'active', 'y'}:
            return True
        if value in {'no', 'false', '0', 'inactive', 'n'}:
            return False
        raise ValidationError(f'Expected yes/no or active/inactive, got: {raw}')
    return str(raw).strip()


def _identity_query(dataset, values):
    name = values.get('name')
    email = values.get('email') or values.get('contact_email')
    if dataset == 'startup':
        if email:
            return Startup.objects.filter(contact_email__iexact=email)
        return Startup.objects.filter(name__iexact=name)
    if dataset == 'mentor':
        if email:
            return Mentor.objects.filter(email__iexact=email)
        return Mentor.objects.filter(name__iexact=name)
    if email:
        return Investor.objects.filter(email__iexact=email)
    return Investor.objects.filter(name__iexact=name, organization__iexact=values.get('organization', ''))


def _identity_key(dataset, values):
    email = (values.get('email') or values.get('contact_email') or '').strip().casefold()
    if email:
        return dataset, 'email', email
    name = _normalize(values.get('name'))
    organization = _normalize(values.get('organization')) if dataset == 'investor' else ''
    return dataset, 'name', name, organization


def analyze_import_preview(batch, mapping):
    """Validate mapped rows without saving and classify likely creates/updates."""
    headers, rows = read_import_file(batch.file.path)
    if batch.dataset == 'mentor':
        normalized = {_normalize(header): header for header in headers}
        first_header, last_header = normalized.get('first name'), normalized.get('last name')
        if first_header and last_header and not any(_normalize(header) in {'name', 'full name', 'mentor name'} for header in headers):
            combined = 'Combined full name'
            headers.append(combined)
            for row in rows:
                row[combined] = ' '.join(str(row.get(key) or '').strip() for key in (first_header, last_header)).strip()

    selected = {field: header for field, header in mapping.items()
                if field in IMPORT_FIELDS[batch.dataset] and header in headers}
    if 'name' not in selected:
        return {'total': len(rows), 'creates': 0, 'updates': 0, 'issues': 1,
                'rows': [{'row': 1, 'status': 'issue', 'message': 'Map a file column to the required Name field.'}]}

    seen = set()
    creates = updates = issues = 0
    details = []
    model = IMPORT_MODELS[batch.dataset]
    for line, row in enumerate(rows, start=2):
        try:
            values = {field: value for field, header in selected.items()
                      if (value := _value_for_field(row.get(header), field)) is not None}
            if not values.get('name'):
                raise ValidationError('Name is required.')
            key = _identity_key(batch.dataset, values)
            if key in seen:
                raise ValidationError('Duplicate identity in this file; only the first matching row can be imported.')
            seen.add(key)
            matches = _identity_query(batch.dataset, values)
            if matches.count() > 1:
                raise ValidationError('This identity matches multiple existing records; resolve duplicates first.')
            instance = matches.first() or model()
            for field, value in values.items():
                setattr(instance, field, value)
            if 'source' in IMPORT_FIELDS[batch.dataset] and not values.get('source'):
                instance.source = batch.original_filename
            instance.full_clean(exclude=['slug'] if batch.dataset == 'startup' else None)
            status = 'update' if instance.pk else 'create'
            if status == 'update':
                updates += 1
            else:
                creates += 1
            message = 'Would update an existing record.' if status == 'update' else 'Would create a new record.'
        except (ValidationError, InvalidOperation, ValueError, TypeError) as exc:
            issues += 1
            status = 'issue'
            message = '; '.join(exc.messages) if isinstance(exc, ValidationError) else str(exc)
        if status == 'issue':
            details.append({'row': line, 'status': status, 'message': message})
    return {'total': len(rows), 'creates': creates, 'updates': updates, 'issues': issues, 'rows': details}


def preview_import(batch):
    headers, rows = read_import_file(batch.file.path)
    if batch.dataset == 'mentor':
        normalized = {_normalize(header): header for header in headers}
        first_header = normalized.get('first name')
        last_header = normalized.get('last name')
        if first_header and last_header and not any(_normalize(header) in {'name', 'full name', 'mentor name'} for header in headers):
            combined = 'Combined full name'
            headers.append(combined)
            for row in rows:
                row[combined] = ' '.join(str(row.get(key) or '').strip() for key in (first_header, last_header)).strip()
    batch.headers = headers
    batch.total_rows = len(rows)
    if not batch.field_mapping:
        batch.field_mapping = guess_field_mapping(batch.dataset, headers)
    batch.save(update_fields=['headers', 'total_rows', 'field_mapping'])
    return headers, rows


def commit_import(batch, mapping):
    if batch.status == 'completed':
        raise ValidationError('This batch has already been imported.')
    headers, rows = read_import_file(batch.file.path)
    if batch.dataset == 'mentor':
        normalized = {_normalize(header): header for header in headers}
        first_header = normalized.get('first name')
        last_header = normalized.get('last name')
        if first_header and last_header and not any(_normalize(header) in {'name', 'full name', 'mentor name'} for header in headers):
            combined = 'Combined full name'
            headers.append(combined)
            for row in rows:
                row[combined] = ' '.join(str(row.get(key) or '').strip() for key in (first_header, last_header)).strip()
    model = IMPORT_MODELS[batch.dataset]
    allowed = set(IMPORT_FIELDS[batch.dataset])
    selected = {field: header for field, header in mapping.items() if field in allowed and header in headers}
    if 'name' not in selected:
        raise ValidationError('Map a file column to the required Name field before importing.')

    created = updated = skipped = 0
    errors = []
    seen = set()
    with transaction.atomic():
        for line, row in enumerate(rows, start=2):
            values = {}
            try:
                for field, header in selected.items():
                    value = _value_for_field(row.get(header), field)
                    if value is not None:
                        values[field] = value
                if not values.get('name'):
                    skipped += 1
                    errors.append({'row': line, 'error': 'Name is required.'})
                    continue
                identity = _identity_key(batch.dataset, values)
                if identity in seen:
                    skipped += 1
                    errors.append({'row': line, 'error': 'Duplicate identity in this file; only the first matching row was imported.'})
                    continue
                seen.add(identity)
                existing = _identity_query(batch.dataset, values)
                if existing.count() > 1:
                    skipped += 1
                    errors.append({'row': line, 'error': 'This identity matches multiple existing records; resolve duplicates before importing.'})
                    continue
                instance = existing.first() or model()
                was_created = instance.pk is None
                for field, value in values.items():
                    setattr(instance, field, value)
                if 'source' in allowed and not values.get('source'):
                    instance.source = batch.original_filename
                instance.full_clean(exclude=['slug'] if batch.dataset == 'startup' else None)
                instance.save()
                if was_created:
                    created += 1
                else:
                    updated += 1
            except (ValidationError, InvalidOperation, ValueError, TypeError) as exc:
                skipped += 1
                message = '; '.join(exc.messages) if isinstance(exc, ValidationError) else str(exc)
                errors.append({'row': line, 'error': message})

    batch.status = 'completed'
    batch.field_mapping = selected
    batch.total_rows = len(rows)
    batch.created_count = created
    batch.updated_count = updated
    batch.skipped_count = skipped
    batch.row_errors = errors
    batch.completed_at = timezone.now()
    batch.save(update_fields=[
        'status', 'field_mapping', 'total_rows', 'created_count', 'updated_count',
        'skipped_count', 'row_errors', 'completed_at',
    ])
    return batch


def report_data(start_date, end_date):
    starts = timezone.make_aware(datetime.combine(start_date, datetime.min.time()))
    ends = timezone.make_aware(datetime.combine(end_date, datetime.max.time()))
    status_rows = list(StartupStatusHistory.objects.filter(recorded_at__range=(starts, ends)).values(
        'status', 'contract_status'
    ).annotate(total=Count('id')).order_by('status', 'contract_status'))
    funding_rows = list(Funding.objects.filter(created_at__range=(starts, ends)).values(
        'currency', 'status', 'investor__name', 'source', 'startup__name'
    ).annotate(total_amount=Sum('amount'), records=Count('id')).order_by('currency', 'status'))
    mentor_rows = list(MentorEngagement.objects.filter(date__range=(start_date, end_date)).values(
        'mentor__name', 'startup__name', 'date', 'hours', 'topics', 'outcome'
    ).order_by('-date'))
    visit_rows = list(PageVisit.objects.filter(visited_at__range=(starts, ends)).values(
        'page_type', 'display_name'
    ).annotate(visits=Count('id')).order_by('-visits', 'display_name'))
    funding_totals = {}
    for row in funding_rows:
        funding_totals[row['currency']] = funding_totals.get(row['currency'], Decimal('0')) + (row['total_amount'] or Decimal('0'))
    duration = (end_date - start_date).days + 1
    previous_end = start_date - timedelta(days=1)
    previous_start = previous_end - timedelta(days=duration - 1)
    previous_start_dt = timezone.make_aware(datetime.combine(previous_start, datetime.min.time()))
    previous_end_dt = timezone.make_aware(datetime.combine(previous_end, datetime.max.time()))
    previous_startups = Startup.objects.filter(created_at__range=(previous_start_dt, previous_end_dt)).count()
    new_startups = Startup.objects.filter(created_at__range=(starts, ends)).count()
    if new_startups > previous_startups:
        activity_signal = f'Startup registrations increased by {new_startups - previous_startups} compared with the previous equal-length period.'
    elif new_startups < previous_startups:
        activity_signal = f'Startup registrations decreased by {previous_startups - new_startups} compared with the previous equal-length period.'
    else:
        activity_signal = 'Startup registrations were unchanged compared with the previous equal-length period.'
    return {
        'start_date': start_date,
        'end_date': end_date,
        'new_startups': new_startups,
        'previous_startups': previous_startups,
        'activity_signal': activity_signal,
        'startups_by_status': list(Startup.objects.values('status').annotate(total=Count('id')).order_by('status')),
        'status_changes': status_rows,
        'new_mentors': Mentor.objects.filter(created_at__range=(starts, ends)).count(),
        'new_investors': Investor.objects.filter(created_at__range=(starts, ends)).count(),
        'funding': funding_rows,
        'funding_totals': funding_totals,
        'kpi_observations': KPI.objects.filter(date_recorded__range=(start_date, end_date)).count(),
        'mentor_sessions': len(mentor_rows),
        'mentor_hours': sum((row['hours'] or Decimal('0')) for row in mentor_rows),
        'mentor_rows': mentor_rows,
        'page_visits': visit_rows,
        'page_visit_total': sum(row['visits'] for row in visit_rows),
    }
